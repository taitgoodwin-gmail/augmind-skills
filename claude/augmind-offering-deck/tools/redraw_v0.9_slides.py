"""Route test only: replay chosen slides of the v0.9 deck data (Google Slides API JSON)
into a .pptx, so we can see whether a code-built file survives conversion to Google Slides.

Usage: replay_to_pptx.py <deck.json> <out.pptx> <slideId> [<slideId> ...]
Not the renderer: it copies stored geometry; the renderer will compute geometry from a spec.
"""
import json
import sys
import copy
from lxml import etree
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.enum.dml import MSO_LINE
from pptx.oxml.ns import qn

HLINK_RGB = RGBColor(0x11, 0x55, 0xCC)
TAB_ROW_Y = 274320          # y of the phase-tab row in v0.9
CARD_RADIUS_EMU = 76200     # 6pt; the API does not expose corner radius, so this is an assumption


def rgb(color):
    if not color:
        return None
    if "themeColor" in color:
        return HLINK_RGB if color["themeColor"] == "HYPERLINK" else RGBColor(0x14, 0x16, 0x1F)
    c = color.get("rgbColor", {})
    return RGBColor(*(round(c.get(k, 0) * 255) for k in ("red", "green", "blue")))


def box(el):
    size, t = el.get("size", {}), el.get("transform", {})
    w = size.get("width", {}).get("magnitude", 0) * t.get("scaleX", 1)
    h = size.get("height", {}).get("magnitude", 0) * t.get("scaleY", 1)
    return Emu(round(t.get("translateX", 0))), Emu(round(t.get("translateY", 0))), Emu(round(w)), Emu(round(h))


def set_bullet(paragraph, glyph, indent_start_pt, indent_first_pt):
    pPr = paragraph._p.get_or_add_pPr()
    pPr.set("marL", str(int(Pt(indent_start_pt))))
    pPr.set("indent", str(int(Pt(indent_first_pt - indent_start_pt))))
    for tag in ("a:buNone", "a:buChar", "a:buAutoNum"):
        for old in pPr.findall(qn(tag)):
            pPr.remove(old)
    bu = etree.SubElement(pPr, qn("a:buChar"))
    bu.set("char", glyph)


def fill_text(shape, text, anchor):
    tf = shape.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = {"TOP": MSO_ANCHOR.TOP, "BOTTOM": MSO_ANCHOR.BOTTOM}.get(anchor, MSO_ANCHOR.MIDDLE)
    para, first, last_size = None, True, None
    for te in text.get("textElements", []):
        if "paragraphMarker" in te:
            para = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            pm = te["paragraphMarker"]
            st = pm.get("style", {})
            para.alignment = {"CENTER": PP_ALIGN.CENTER, "END": PP_ALIGN.RIGHT}.get(st.get("alignment"), PP_ALIGN.LEFT)
            para.line_spacing = st.get("lineSpacing", 100) / 100.0
            para.space_before = Pt(st.get("spaceAbove", {}).get("magnitude", 0))
            para.space_after = Pt(st.get("spaceBelow", {}).get("magnitude", 0))
            if "bullet" in pm:
                set_bullet(para, pm["bullet"].get("glyph", "•"),
                           st.get("indentStart", {}).get("magnitude", 12),
                           st.get("indentFirstLine", {}).get("magnitude", 0))
        elif "textRun" in te and para is not None:
            content = te["textRun"]["content"].rstrip("\n")
            st = te["textRun"].get("style", {})
            if not content:
                continue
            for i, piece in enumerate(content.split("\u000b")):
                if i:
                    para.add_line_break()
                if not piece:
                    continue
                run = para.add_run()
                run.text = piece
                f = run.font
                f.name = st.get("fontFamily", "Inter")
                if "fontSize" in st:
                    f.size = Pt(st["fontSize"]["magnitude"])
                    last_size = f.size
                f.bold = bool(st.get("bold"))
                f.italic = bool(st.get("italic"))
                url = st.get("link", {}).get("url")
                if url:
                    run.hyperlink.address = url
                col = rgb(st.get("foregroundColor", {}).get("opaqueColor"))
                if col is not None:
                    f.color.rgb = col
                if st.get("underline") or url:
                    f.underline = True
    return last_size


def build(deck_path, out_path, slide_ids):
    deck = json.load(open(deck_path))["content"]["slides"]
    by_id = {s["objectId"]: s for s in deck}
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
    blank = prs.slide_layouts[6]
    made = {sid: prs.slides.add_slide(blank) for sid in slide_ids}
    stats = {"shapes": 0, "url_links": 0, "slide_links": 0, "slide_links_skipped": 0, "tables_skipped": 0}
    for sid in slide_ids:
        src, slide = by_id[sid], made[sid]
        bg = src.get("pageProperties", {}).get("pageBackgroundFill", {}).get("solidFill", {}).get("color")
        if bg:
            slide.background.fill.solid()
            slide.background.fill.fore_color.rgb = rgb(bg)
        is_cover = sid == slide_ids[0] and sid == "p1"
        for el in src["pageElements"]:
            if "table" in el:
                stats["tables_skipped"] += 1
                continue
            sh = el.get("shape")
            if not sh:
                continue
            left, top, w, h = box(el)
            props = sh.get("shapeProperties", {})
            kind = sh.get("shapeType")
            if kind == "TEXT_BOX":
                shape = slide.shapes.add_textbox(left, top, w, h)
            else:
                auto = MSO_SHAPE.ROUNDED_RECTANGLE if kind == "ROUND_RECTANGLE" else MSO_SHAPE.RECTANGLE
                shape = slide.shapes.add_shape(auto, left, top, w, h)
                shape.shadow.inherit = False
                style = shape._element.find(qn("p:style"))   # drop theme style so no shadow is inherited
                if style is not None:
                    shape._element.remove(style)
                if kind == "ROUND_RECTANGLE":
                    pill = is_cover or abs(int(top) - TAB_ROW_Y) < 2
                    shape.adjustments[0] = 0.5 if pill else min(0.5, CARD_RADIUS_EMU / max(1, min(int(w), int(h))))
            fill = props.get("shapeBackgroundFill", {})
            if fill.get("propertyState") != "NOT_RENDERED" and "solidFill" in fill:
                shape.fill.solid()
                shape.fill.fore_color.rgb = rgb(fill["solidFill"].get("color"))
            else:
                shape.fill.background()
            line = props.get("outline", {})
            if kind != "TEXT_BOX" and line.get("propertyState") != "NOT_RENDERED" and line.get("outlineFill"):
                shape.line.color.rgb = rgb(line["outlineFill"]["solidFill"].get("color"))
                shape.line.width = Emu(int(line.get("weight", {}).get("magnitude", 9525)))
                if line.get("dashStyle") == "DASH":
                    shape.line.dash_style = MSO_LINE.DASH
            else:
                shape.line.fill.background()
            if "text" in sh:
                fill_text(shape, sh["text"], props.get("contentAlignment"))
                stats["url_links"] += sum(1 for te in sh["text"].get("textElements", [])
                                          if te.get("textRun", {}).get("style", {}).get("link", {}).get("url"))
            target = props.get("link", {}).get("pageObjectId")
            if target:
                if target in made:
                    shape.click_action.target_slide = made[target]
                    stats["slide_links"] += 1
                else:
                    stats["slide_links_skipped"] += 1
            stats["shapes"] += 1
    prs.save(out_path)
    return stats


if __name__ == "__main__":
    print(build(sys.argv[1], sys.argv[2], sys.argv[3:]))
