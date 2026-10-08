"""Compare the regenerated deck with v0.9 slide by slide: every word and every link."""
import json, re, collections, sys
from pptx import Presentation
deck = json.load(open(sys.argv[1]))["content"]["slides"]
prs = Presentation(sys.argv[2])
def words(t): return re.findall(r"[^\s·›•]+", t)
def ref(s):
    w, urls = [], []
    def take(tes):
        w.extend(words("".join(te["textRun"]["content"] for te in tes if "textRun" in te)))   # join runs: v0.9 splits some words across runs
        for te in tes:
            u = te.get("textRun", {}).get("style", {}).get("link", {}).get("url")
            if u: urls.append(u)
    for e in s["pageElements"]:
        if "table" in e:
            for r in e["table"]["tableRows"]:
                for c in r["tableCells"]: take(c.get("text", {}).get("textElements", []))
        else: take(e["shape"].get("text", {}).get("textElements", []))
    return collections.Counter(w), sorted(urls)
bad = 0
for i, (slide, s) in enumerate(zip(prs.slides, deck), 1):
    w, urls = [], []
    for sh in slide.shapes:
        if sh.has_text_frame:
            for para in sh.text_frame.paragraphs:
                w.extend(words("".join(r.text for r in para.runs)))
                for r in para.runs:
                    if r.hyperlink.address: urls.append(r.hyperlink.address)
    rw, ru = ref(s); mw = collections.Counter(w)
    missing, extra = rw - mw, mw - rw
    same_links = ru == sorted(urls)
    if missing or extra or not same_links:
        bad += 1
        print(f"slide {i:2d} ({s['objectId']}): missing {dict(missing) or '-'} | extra {dict(extra) or '-'} | links {len(ru)} vs {len(urls)} {'same' if same_links else 'DIFFERENT'}")
print(f"slides compared: {min(len(prs.slides), len(deck))} of v0.9's {len(deck)} (deck has {len(prs.slides)}); slides with any difference: {bad}")
